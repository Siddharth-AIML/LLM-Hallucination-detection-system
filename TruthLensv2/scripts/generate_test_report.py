
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent
REPORTS = ROOT / "reports"
REPORTS.mkdir(exist_ok=True)

# Table 6.1: Automated Test Results
suite_rows = [
    ["Context Module Contract Tests", "test_context_modules_contract.py",
     8, 8, "43.04 s*", "Context audit, memory analysis, workflow error handling"],
    ["Extreme Input Tests", "test_extreme_inputs.py",
     7, 7, "Included in full-suite run", "Oversized inputs, Unicode, multiline text"],
    ["Factual Pipeline Contract Tests", "test_factual_pipeline_contract.py",
     3, 3, "45.78 s*", "Factual result schema, structured errors, empty output"],
    ["GSM8K Benchmark Tests", "test_gsm8k_benchmark.py",
     2, 2, "Included in full-suite run", "Numeric answer extraction and benchmark metrics"],
    ["Live Functionality Tests", "test_live_functionality_optional.py",
     3, 0, "Skipped", "Real configured modules; requires TRUTHLENS_RUN_LIVE=1"],
    ["Math API Tests", "test_math_api.py",
     5, 5, "Included in full-suite run", "Math API validation and route availability"],
    ["Math Verifier Tests", "test_math_verifier.py",
     20, 20, "Included in full-suite run", "Arithmetic, word problems, ambiguous cases"],
    ["Request Validation Tests", "test_request_validation.py",
     14, 14, "Included in full-suite run", "Required fields, sizes, types, health endpoint"],
    ["Full Suite Summary", "All test modules",
     62, "59 passed", "44.07 s", "3 skipped, 0 failures, 4 warnings"],
]

# Table 6.2: Edge-case tests
factual = [
    ("F01", "Supported claim", "Supporting evidence and verdict are serialized", "Pass"),
    ("F02", "Contradicted claim", "Contradicting evidence and verdict are serialized", "Pass"),
    ("F03", "No evidence available", "Claim is not promoted to supported without evidence", "Pass"),
    ("F04", "Multiple claims", "Each claim retains its own verdict", "Pass"),
    ("F05", "Conflicting evidence", "Both supporting and contradicting evidence are preserved", "Pass"),
    ("F06", "Date and numeric claims", "Per-claim results preserve response structure", "Pass"),
    ("F07", "Duplicate claims", "Duplicate claims do not break response structure", "Pass"),
    ("F08", "Empty input", "A safe empty result is returned", "Pass"),
    ("F09", "Unicode, multiline, injection-like text", "Input is forwarded as data to the simulated pipeline", "Pass"),
    ("F10", "Pipeline failure", "A structured error response is returned", "Pass"),
]
contextual = [
    ("C01", "Context contradiction", "Contradiction finding is returned", "Pass"),
    ("C02", "No-action mitigation", "Supported answer does not trigger unnecessary mitigation", "Pass"),
    ("C03", "Missing context", "Insufficient-context result is returned by the simulated auditor", "Pass"),
    ("C04", "Instruction violation", "Instruction-violation finding is preserved", "Pass"),
    ("C05", "Irrelevant answer", "Answer is passed for alignment auditing", "Pass"),
    ("C06", "Empty conversation history", "Memory analysis is marked not assessable", "Pass"),
    ("C07", "Conflicting history and injection-like text", "Inputs are passed as untrusted data", "Pass"),
    ("C08", "Partial workflow failure", "Workflow reports partial failure safely", "Pass"),
    ("C09", "Mitigation re-verification", "Remaining issue status is returned", "Pass"),
    ("C10", "Validation/backend failure", "Validation and simulated backend errors are handled", "Pass"),
]

# ---------- Excel report ----------
wb = Workbook()
ws = wb.active
ws.title = "Automated Test Results"

ws.append(["Table 6.1: Automated Test Results"])
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
ws.append([
    "Test Suite Package", "Test Module File", "Tests Collected",
    "Tests Passed", "Execution Time", "Primary Subsystem Evaluated"
])
for row in suite_rows:
    ws.append(row)

edge = wb.create_sheet("Edge-Case Results")
edge.append(["Table 6.2: Edge-Case Test Results"])
edge.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
edge.append([
    "Case ID", "Test Module", "Input / Test Scenario",
    "Expected / Observed Behaviour", "Result"
])
for case, scenario, behaviour, result in factual:
    edge.append([case, "Factual", scenario, behaviour, result])
for case, scenario, behaviour, result in contextual:
    edge.append([case, "Contextual", scenario, behaviour, result])
edge.append(["Total", "Both suites", "20 edge-case tests", "20 passed; 0 failed", "Pass"])

thin = Side(style="thin", color="000000")
for sheet in wb.worksheets:
    sheet.freeze_panes = "A3"
    sheet.sheet_view.showGridLines = False
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(
                vertical="center", horizontal="left", wrap_text=True
            )
            cell.font = Font(name="Times New Roman", size=10)
    for cell in sheet[1]:
        cell.font = Font(name="Times New Roman", size=12, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for cell in sheet[2]:
        cell.font = Font(name="Times New Roman", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 25
    sheet.row_dimensions[2].height = 35

for col, width in enumerate([28, 34, 16, 16, 22, 48], 1):
    ws.column_dimensions[get_column_letter(col)].width = width
for col, width in enumerate([12, 16, 31, 55, 12], 1):
    edge.column_dimensions[get_column_letter(col)].width = width

for sheet in wb.worksheets:
    for row in range(3, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 36

xlsx_path = REPORTS / "TruthLensV2_Test_Results.xlsx"
wb.save(xlsx_path)

# ---------- Word report with black-book-style tables ----------
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.55)
section.right_margin = Inches(0.55)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(9)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("TruthLens V2 – Automated Test Results")
r.bold = True
r.font.name = "Times New Roman"
r.font.size = Pt(13)

def add_word_table(caption, headers, rows):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(caption)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True

    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)

    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = 1
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                for run in para.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8)
            if ri == 0:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.bold = True

add_word_table(
    "Table 6.1: Automated Test Results",
    ["Test Suite Package", "Test Module File", "Tests Collected",
     "Tests Passed", "Execution Time", "Primary Subsystem Evaluated"],
    suite_rows
)

doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("TruthLens V2 – Edge-Case Test Results")
r.bold = True
r.font.name = "Times New Roman"
r.font.size = Pt(13)

edge_rows = []
for case, scenario, behaviour, result in factual:
    edge_rows.append([case, "Factual", scenario, behaviour, result])
for case, scenario, behaviour, result in contextual:
    edge_rows.append([case, "Contextual", scenario, behaviour, result])
edge_rows.append(["Total", "Both suites", "20 edge-case tests", "20 passed; 0 failed", "Pass"])

add_word_table(
    "Table 6.2: Edge-Case Test Results",
    ["Case ID", "Test Module", "Input / Test Scenario",
     "Expected / Observed Behaviour", "Result"],
    edge_rows
)

doc.add_paragraph()
p = doc.add_paragraph()
run = p.add_run("Testing limitation: ")
run.bold = True
p.add_run(
    "Several edge-case tests use mocked/simulated pipeline or module responses. "
    "These results validate API contracts, response serialization, input validation, "
    "and error handling; they do not independently prove live LLM hallucination-detection accuracy."
)

docx_path = REPORTS / "TruthLensV2_Test_Results.docx"
doc.save(docx_path)

print("Reports generated successfully:")
print(xlsx_path)
print(docx_path)
